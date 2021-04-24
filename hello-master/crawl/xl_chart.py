import openpyxl
from openpyxl.chart import (
    Reference,
    BarChart
)

filename = './data/output.xlsx'
book = openpyxl.load_workbook(filename)
sheet = book.create_sheet()

rows = [
    ['김일수', 11],
    ['김이수', 22],
    ['김삼수', 33],
    ['김사수', 15],
    ['김오수', 11],
]

for row in rows:
    sheet.append(row)

datax = Reference(sheet, min_col=2,
                  min_row=1, max_col=2, max_row=5)
categs = Reference(sheet, min_col=1,
                   min_row=1, max_row=5)

chart = BarChart()
chart.add_data(data=datax)
chart.set_categories(categs)

chart.legend = None  # 범례
chart.varyColors = True
chart.title = "차트 타이틀"

sheet.add_chart(chart, "A8")

book.save(filename)
