from PIL import Image
import openpyxl
from pprint import pprint
from openpyxl.chart import (
    Reference,
    BarChart
)

file = "./t1.xlsx"
book = openpyxl.load_workbook(file)
# sheet=book.worksheets[0]
sheet = book.create_sheet()
sheet.title = "이미지테스트"

imgFile='../crawl/images/aaa.png'
img=openpyxl.drawing.image.Image(imgFile)
sheet.add_image(img, 'B5')

img2=Image.open(imgFile)
new_img=img2.resize((30, 30))
new_img.save('aaa.png')

img3=openpyxl.drawing.image.Image('aaa.png')
sheet.add_image(img3, 'D1')
img4 = openpyxl.drawing.image.Image('aaa.png')
sheet.add_image(img4, 'D2')


# chart
sheet1 = book.worksheets[0]
data = Reference(sheet1, min_col=5, min_row=2, max_col=5, max_row=11)
cates = Reference(sheet1, min_col=1, min_row=2, max_row=11)

chart = BarChart()
chart.add_data(data=data)
chart.set_categories(cates)

chart.legend = 'Legend'
# chart.y_axis.majorGridlines = None
chart.varyColors = True
chart.title = "Top 10 Likes"

sheet.add_chart(chart, "A8")

book.save(file)
