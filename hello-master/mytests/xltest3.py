from PIL import Image
import openpyxl
from pprint import pprint
from openpyxl.chart import (
    Reference,
    BarChart
)

file = "./t1.xlsx"
book = openpyxl.load_workbook(file)
sheet = book.create_sheet()
sheet.title = "Chart테스트"

# rows = [
#     ("USA", 46),
#     ("China", 38),
#     ("UK", 29),
#     ("South Korea", 13),
#     ("Germany", 11)
# ]

rows = [
    ['김일수', 11],
    ['김이수', 22],
    ['김삼수', 33],
    ['김사수', 15],
    ['김오수', 11],
]

for row in rows:
    sheet.append(row)

data = Reference(sheet, min_col=2, min_row=1, max_col=2, max_row=5)
categs = Reference(sheet, min_col=1, min_row=1, max_row=5)

chart = BarChart()
chart.add_data(data=data)
chart.set_categories(categs)

chart.legend = None  # 범례
chart.varyColors = True
chart.title = "차트 타이틀"

sheet.add_chart(chart, "A8")

book.save(file)
